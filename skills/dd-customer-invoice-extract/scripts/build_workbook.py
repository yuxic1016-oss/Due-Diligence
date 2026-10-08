#!/usr/bin/env python3
"""
Build the diligence workbook from extract_invoices.py output.

Usage:
    python build_workbook.py extracted.json -o "Deal Invoice Extract.xlsx"
    python build_workbook.py extracted.json -o out.xlsx --mrc-scope all

Tabs: Summary | Line Items | Charge Categories | Assumptions & Notes

Every derived number is a live formula referencing the Line Items tab, and the
derate sits in a single cell so a reviewer who disagrees with 80% can reprice the
whole book by editing one number.

After running this, recalculate the file (LibreOffice via the xlsx skill's
recalc.py) and confirm zero formula errors -- openpyxl writes formulas with no
cached values, so the file reads back empty until it is recalculated.
"""
import argparse, json, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gcl

CATS = ["Power", "Space / Cabinet", "Cross Connect / Fiber", "Bandwidth", "Ports",
        "IP Addresses", "Rent / CAM / Utilities", "Chilled Water", "Remote Hands",
        "Hardware / PDU", "Master Service ($0)", "Setup Fee (one-time)",
        "Late Fee (one-time)", "Other"]

CAL = lambda **k: Font(name="Calibri", size=11, **k)
BLUE, BLACK, GREEN = CAL(color="0000FF"), CAL(), CAL(color="008000")
BLACK_B = CAL(bold=True)
HDRF = CAL(bold=True, color="FFFFFF")
HDRFILL = PatternFill("solid", fgColor="000000")
KEYFILL = PatternFill("solid", fgColor="FFF2CC")
TOPB = Border(top=Side(style="thin"))
CUR = "$#,##0.00;($#,##0.00);-"
NUM = "#,##0.00;(#,##0.00);-"


def header(ws, row, labels):
    for j, h in enumerate(labels, 1):
        c = ws.cell(row, j, h)
        c.font, c.fill = HDRF, HDRFILL
        c.alignment = Alignment(horizontal="center", wrap_text=True, vertical="center")


def build(data, out, mrc_scope):
    meta = data["meta"]
    recs = data["invoices"]
    sites = {}
    for r in recs:
        sites.setdefault(r["site"], len(sites))
    recs.sort(key=lambda r: (sites[r["site"]], r.get("inv", "")))
    red_label = ("2N — A and B feeds both counted" if meta["redundancy"] == "2N"
                 else "N — one feed of an A+B pair counted")

    wb = Workbook()

    # ------------------------------------------------ Assumptions & Notes
    asm = wb.active
    asm.title = "Assumptions & Notes"
    asm.showGridLines = False
    flagged = [(r.get("inv", r["file"]), i["desc"], i["kwd"]["flag"])
               for r in recs for i in r["items"] if i.get("kwd", {}).get("flag")]
    onetime_total = sum(i["total"] or 0 for r in recs for i in r["items"] if i["onetime"])
    onetime_count = len({r["file"] for r in recs
                         for i in r["items"] if i["onetime"]})

    rows = [
        ("Customer Invoice Extract", "", ""),
        (f"Source: {meta['folder']}  ({meta['invoice_count']} invoice PDFs)", "", ""),
        ("", "", ""),
        ("kW derivation assumptions", "", ""),
        ("Derate factor applied to breaker rating", "", meta["derate"]),
        ("Redundancy treatment", red_label, ""),
        ("Single-phase kW formula", "Volts x Amps x Circuits x Derate x Redundancy / 1000", ""),
        ("Three-phase multiplier", "sqrt(3)", 1.7320508),
        ("", "", ""),
        ("Notes on the kW figure", "", ""),
        ("kW is not stated on these invoices. It is derived from the billed power circuits (volts x amps).", "", ""),
        ("Where the quantity column equals the stated amperage, the provider bills per amp: circuits = 1, amps = quantity.", "", ""),
        ("Where a SKU code and the description text disagree on amperage, the SKU is used (flagged per line below).", "", ""),
        ("Plug-access, PDU rental and hardware lines are excluded from kW — they are not billed capacity.", "", ""),
        ("Wholesale rent / CAM tenants carry no derivable kW; power reaches them as an estimated electric charge inside rent.", "", ""),
        ("", "", ""),
        ("MRC scope", "", ""),
        (("Total MRC = all service-item lines excluding one-time charges (late fees, setup/install). Pre-tax."
          if mrc_scope == "recurring" else
          "Total MRC = the Service Items subtotal printed on each invoice. Pre-tax, all lines included."), "", ""),
        (f"One-time charges identified: ${onetime_total:,.2f} across {onetime_count} invoice(s).", "", ""),
        ("", "", ""),
        ("Reconciliation", "", ""),
        (f"Line items tie to the printed Service Items subtotal on "
         f"{meta['reconciled']} of {meta['invoice_count']} invoices.", "", ""),
    ]
    if flagged:
        rows += [("", "", ""), ("Lines carrying an estimate or a resolved conflict", "", "")]
        rows += [(f"{inv} — {d[:90]}", f, "") for inv, d, f in flagged]

    bold_rows = {1, 4, 10, 17, 21}
    for i, (a, b, c) in enumerate(rows, 1):
        cell = asm.cell(i, 1, a)
        cell.font = BLACK_B if i in bold_rows or a.endswith("estimate or a resolved conflict") else BLACK
        if b:
            asm.cell(i, 2, b).font = BLACK
        if c != "":
            asm.cell(i, 3, c).font = BLUE
    asm.cell(5, 3).number_format = "0%"
    asm.cell(8, 3).number_format = "0.0000000"
    for col, w in (("A", 118), ("B", 48), ("C", 14)):
        asm.column_dimensions[col].width = w

    # --------------------------------------------------------- Line Items
    li = wb.create_sheet("Line Items")
    li.showGridLines = False
    header(li, 1, ["Site", "Invoice #", "Customer", "Invoice Date", "Item ID",
                   "Description", "Charge Category", "Code", "Service Period Start",
                   "Unit Price", "Qty", "Line Amount", "Recurring?", "Volts", "Amps",
                   "Circuits", "Phase Factor", "Redundancy (N)", "Derived kW", "Note"])
    li.freeze_panes = "C2"
    r = 1
    for rec in recs:
        for it in rec["items"]:
            r += 1
            k = it.get("kwd") or {}
            vals = [rec["site"], rec.get("inv", ""), rec.get("cust", ""),
                    rec.get("date", ""), it["id"], it["desc"], it["cat"], it["code"],
                    it.get("range_start", ""), it["unit"], it["qty"], it["total"],
                    "One-time" if it["onetime"] else "Recurring",
                    k.get("V", ""), k.get("A", ""), k.get("circ", ""),
                    k.get("ph", ""), k.get("red", ""), None, k.get("flag", "")]
            for j, v in enumerate(vals, 1):
                li.cell(r, j, v).font = BLACK if j in (7, 13, 19) else BLUE
            if k and k.get("V") not in ("", None):
                li.cell(r, 19).value = (f"=N{r}*O{r}*P{r}*Q{r}*R{r}"
                                        f"*'Assumptions & Notes'!$C$5/1000")
                li.cell(r, 19).font = GREEN
            else:
                li.cell(r, 19).value = it["kw"]
                li.cell(r, 19).font = BLUE
            for j in (10, 12):
                li.cell(r, j).number_format = CUR
            for j in (11, 15, 16, 18, 19):
                li.cell(r, j).number_format = NUM
    last_li = r
    li.cell(r + 1, 6, "TOTAL").font = BLACK_B
    for j, col in ((12, "L"), (19, "S")):
        c = li.cell(r + 1, j, f"=SUM({col}2:{col}{r})")
        c.font, c.number_format = BLACK_B, (CUR if j == 12 else NUM)
    for j in range(1, 21):
        li.cell(r + 1, j).border = TOPB
    for col, w in zip("ABCDEFGHIJKLMNOPQRST",
                      [16, 16, 32, 13, 9, 72, 22, 13, 17, 12, 9, 13, 12, 8, 8, 9, 11, 13, 11, 52]):
        li.column_dimensions[col].width = w

    # ------------------------------------------------------------ Summary
    sm = wb.create_sheet("Summary", 0)
    sm.showGridLines = False
    header(sm, 1, ["Site", "Invoice Date", "Customer", "Invoice #", "Derived kW",
                   "Total MRC", "One-Time Charges", "Service Items Subtotal", "Tax",
                   "Invoice Total"] + CATS)
    sm.freeze_panes = "E2"
    LI = "'Line Items'"
    for i, rec in enumerate(recs, 2):
        for j, v in ((1, rec["site"]), (2, rec.get("date", "")),
                     (3, rec.get("cust", "")), (4, rec.get("inv", ""))):
            sm.cell(i, j, v).font = BLUE
        inv = f"$D{i}"
        sumifs = lambda col, extra="": (
            f"=SUMIFS({LI}!${col}$2:${col}${last_li},"
            f"{LI}!$B$2:$B${last_li},{inv}{extra})")
        sm.cell(i, 5, sumifs("S")).font = GREEN
        if mrc_scope == "recurring":
            sm.cell(i, 6, sumifs("L", f',{LI}!$M$2:$M${last_li},"Recurring"')).font = GREEN
            sm.cell(i, 7, sumifs("L", f',{LI}!$M$2:$M${last_li},"One-time"')).font = GREEN
        else:
            sm.cell(i, 6, sumifs("L")).font = GREEN
            sm.cell(i, 7, 0).font = BLUE
        sm.cell(i, 8, f"=F{i}+G{i}").font = BLACK
        tax = round((rec.get("inv_total") or 0) - (rec.get("svc_total") or 0), 2)
        sm.cell(i, 9, tax if abs(tax) > 0.004 else 0).font = BLUE
        sm.cell(i, 10, f"=H{i}+I{i}").font = BLACK
        for jj, _ in enumerate(CATS):
            sm.cell(i, 11 + jj,
                    sumifs("L", f",{LI}!$G$2:$G${last_li},{gcl(11 + jj)}$1")).font = GREEN
        sm.cell(i, 5).number_format = NUM
        for j in range(6, 11 + len(CATS)):
            sm.cell(i, j).number_format = CUR
    last_sm = len(recs) + 1
    tr = last_sm + 1
    sm.cell(tr, 3, f"TOTAL — {len(recs)} invoices").font = BLACK_B
    for j in range(5, 11 + len(CATS)):
        L = gcl(j)
        c = sm.cell(tr, j, f"=SUM({L}2:{L}{last_sm})")
        c.font, c.number_format = BLACK_B, (NUM if j == 5 else CUR)
    for j in range(1, 11 + len(CATS)):
        sm.cell(tr, j).border = TOPB
    sm.cell(tr, 5).fill = KEYFILL
    sm.cell(tr, 6).fill = KEYFILL
    for col, w in zip("ABCDEFGHIJ", [16, 13, 34, 16, 11, 13, 13, 14, 10, 13]):
        sm.column_dimensions[col].width = w
    for jj in range(len(CATS)):
        sm.column_dimensions[gcl(11 + jj)].width = 15

    # -------------------------------------------------- Charge Categories
    cc = wb.create_sheet("Charge Categories")
    cc.showGridLines = False
    header(cc, 1, ["Charge Category", "# of Line Items", "Total Amount",
                   "% of Total", "Derived kW"])
    for i, cat in enumerate(CATS, 2):
        cc.cell(i, 1, cat).font = BLACK
        cc.cell(i, 2, f"=COUNTIF({LI}!$G$2:$G${last_li},$A{i})").font = GREEN
        cc.cell(i, 3, f"=SUMIF({LI}!$G$2:$G${last_li},$A{i},{LI}!$L$2:$L${last_li})").font = GREEN
        cc.cell(i, 4, f"=IFERROR(C{i}/$C${len(CATS) + 2},0)").font = BLACK
        cc.cell(i, 5, f"=SUMIF({LI}!$G$2:$G${last_li},$A{i},{LI}!$S$2:$S${last_li})").font = GREEN
        cc.cell(i, 3).number_format = CUR
        cc.cell(i, 4).number_format = "0.0%"
        cc.cell(i, 5).number_format = NUM
    tr2 = len(CATS) + 2
    cc.cell(tr2, 1, "TOTAL").font = BLACK_B
    for j, fmt in ((2, "#,##0"), (3, CUR), (4, "0.0%"), (5, NUM)):
        L = gcl(j)
        c = cc.cell(tr2, j, f"=SUM({L}2:{L}{tr2 - 1})")
        c.font, c.number_format = BLACK_B, fmt
    for j in range(1, 6):
        cc.cell(tr2, j).border = TOPB
    for col, w in zip("ABCDE", [26, 16, 16, 12, 12]):
        cc.column_dimensions[col].width = w

    wb._sheets = [wb[n] for n in ["Summary", "Line Items", "Charge Categories",
                                  "Assumptions & Notes"]]

    wb.save(out)
    print(f"wrote {out}  ({len(recs)} invoices, {last_li - 1} line items)")
    print("next: recalculate the file and confirm zero formula errors before delivering.")


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("json_file")
    ap.add_argument("-o", "--output", default="Invoice Extract.xlsx")
    ap.add_argument("--mrc-scope", choices=["recurring", "all"], default="recurring",
                    help="recurring excludes late/setup fees; all uses every service item")
    args = ap.parse_args()
    with open(args.json_file) as fh:
        build(json.load(fh), args.output, args.mrc_scope)


if __name__ == "__main__":
    main()
