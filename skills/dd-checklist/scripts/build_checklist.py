#!/usr/bin/env python3
"""Write STATUS and NOTE into a DD checklist workbook, in house format.

Usage:
    python3 build_checklist.py CHECKLIST.xlsx mapping.py -o OUTPUT.xlsx

mapping.py is a plain Python file defining:

    HEADER_ROW = 12          # row holding the checklist's column headers
    SHEET      = "Sheet1"    # source sheet name; renamed to "Checklist" on output

    # row -> (status, note, documents, path, verified_note)
    # status is "Received", "Partial", or "" / None to leave the row blank.
    # documents / path / verified_note are optional and go to the Detail tab.
    ROWS = {
        15:  ("Received", "", "5.1.3.1_Asset Purchase Agreement.docx", "5.1_Transaction Documents", ""),
        31:  ("Partial", "Waiting on benefits", "5.3.5.2.2_Employee Census.xlsx", "5.3.5_Personnel", ""),
        129: ("Partial", "From seller: Cash basis is not available", "", "", ""),
    }

    DATA_ROOM = "Data Room 8.12.26"          # optional, for the summary tab
    TRACKER   = []                            # optional: (item, request, seller_status, verdict, evidence)

House format, per references/output-format.md:
  - working tab named "Checklist", columns A-E untouched
  - F STATUS, G NOTE only; Times New Roman 11, centred, wrapped, NO fill
  - gridlines off; blank rows stay blank
  - detail lives on separate tabs
"""
import argparse
import importlib.util
import os
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BODY = Font(name="Times New Roman", size=11)
HEAD = Font(name="Times New Roman", size=11, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
TOPWRAP = Alignment(vertical="top", wrap_text=True)
MEDIUM = Border(left=Side(style="medium"), right=Side(style="medium"),
                top=Side(style="medium"), bottom=Side(style="medium"))
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name="Times New Roman", size=11, bold=True, color="FFFFFF")


def load_mapping(path):
    spec = importlib.util.spec_from_file_location("mapping", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def norm(entry):
    """Accept 2- to 5-tuples."""
    entry = tuple(entry) + ("",) * (5 - len(entry))
    return entry[:5]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("checklist")
    ap.add_argument("mapping")
    ap.add_argument("-o", "--out")
    a = ap.parse_args()

    m = load_mapping(a.mapping)
    header_row = getattr(m, "HEADER_ROW", 12)
    src_sheet = getattr(m, "SHEET", None)
    data_room = getattr(m, "DATA_ROOM", "")
    tracker = getattr(m, "TRACKER", [])
    rows = {r: norm(v) for r, v in m.ROWS.items()}

    wb = openpyxl.load_workbook(a.checklist)
    ws = wb[src_sheet] if src_sheet else wb.worksheets[0]
    ws.title = "Checklist"

    # ---- headers -------------------------------------------------------
    for col, label in ((6, "STATUS"), (7, "NOTE")):
        c = ws.cell(header_row, col, label)
        c.font = HEAD
        c.alignment = CENTER
        c.border = MEDIUM

    # ---- body ----------------------------------------------------------
    written = 0
    for row, (status, note, _docs, _path, _vnote) in sorted(rows.items()):
        if not status:
            continue                       # blank stays blank
        if status not in ("Received", "Partial"):
            raise SystemExit(
                "row %d: status %r not allowed on the Checklist tab - collapse to "
                "'Received', 'Partial', or blank (see references/status-rules.md)"
                % (row, status))
        f = ws.cell(row, 6, status)
        g = ws.cell(row, 7, note or None)
        for c in (f, g):
            c.font = BODY
            c.alignment = CENTER
        written += 1

    ws.column_dimensions["F"].width = 17
    ws.column_dimensions["G"].width = 32.5
    ws.sheet_view.showGridLines = False

    # ---- Detail tab ----------------------------------------------------
    d = wb.create_sheet("Detail", 1)
    for j, label in enumerate(
            ["Row", "Line", "Item", "Status", "Note",
             "Document(s) received", "Data room path", "Notes from opening the document"], start=1):
        c = d.cell(1, j, label)
        c.font = HDR_FONT
        c.fill = HDR_FILL
        c.alignment = CENTER
    dr = 2
    for row, (status, note, docs, path, vnote) in sorted(rows.items()):
        item = ws.cell(row, 2).value or ws.cell(row, 3).value or ""
        line = ws.cell(row, 1).value
        vals = [row, str(line).lstrip("="), str(item), status or "", note or "",
                docs or "", path or "-", vnote or ""]
        for j, v in enumerate(vals, start=1):
            c = d.cell(dr, j, v)
            c.font = BODY
            c.alignment = TOPWRAP
        dr += 1
    for col, w in zip("ABCDEFGH", (7, 8, 46, 13, 34, 70, 52, 70)):
        d.column_dimensions[col].width = w
    d.freeze_panes = "A2"
    d.auto_filter.ref = "A1:H%d" % (dr - 1)

    # ---- Status Summary ------------------------------------------------
    s = wb.create_sheet("Status Summary", 2)
    s["A1"] = "Checklist status" + (" vs. " + data_room if data_room else "")
    s["A1"].font = Font(name="Times New Roman", size=14, bold=True)
    s["A2"] = "Blank status = not yet received or not applicable. Partial rows carry what is missing in the note."
    s["A2"].font = Font(name="Times New Roman", size=9, italic=True)
    r = 4
    for label in ("Received", "Partial", "Blank / outstanding"):
        if label == "Blank / outstanding":
            n = sum(1 for v in rows.values() if not v[0])
        else:
            n = sum(1 for v in rows.values() if v[0] == label)
        s.cell(r, 1, label).font = HEAD
        s.cell(r, 2, n).font = BODY
        r += 1
    s.cell(r, 1, "Lines mapped").font = HEAD
    s.cell(r, 2, len(rows)).font = HEAD

    r += 2
    s.cell(r, 1, "PARTIAL / DISPUTED - what is outstanding").font = Font(
        name="Times New Roman", size=12, bold=True)
    r += 1
    for j, label in enumerate(["Checklist line", "Item", "What is missing"], start=1):
        c = s.cell(r, j, label)
        c.font = HDR_FONT
        c.fill = HDR_FILL
    r += 1
    for row, (status, note, _d, _p, _v) in sorted(rows.items()):
        if status != "Partial":
            continue
        item = ws.cell(row, 2).value or ws.cell(row, 3).value or ""
        line = str(ws.cell(row, 1).value or "").lstrip("=")
        for j, v in enumerate((line, str(item), note or ""), start=1):
            c = s.cell(r, j, v)
            c.font = BODY
            c.alignment = TOPWRAP
        r += 1
    for col, w in zip("ABC", (16, 52, 90)):
        s.column_dimensions[col].width = w

    # ---- Tracker tab ---------------------------------------------------
    if tracker:
        t = wb.create_sheet("Tracker vs Data Room", 3)
        for j, label in enumerate(
                ["Item #", "Request", "Seller status", "Verdict vs data room", "What I actually see"], start=1):
            c = t.cell(1, j, label)
            c.font = HDR_FONT
            c.fill = HDR_FILL
            c.alignment = CENTER
        for i, entry in enumerate(tracker, start=2):
            for j, v in enumerate(entry, start=1):
                c = t.cell(i, j, v)
                c.font = BODY
                c.alignment = TOPWRAP
        for col, w in zip("ABCDE", (8, 40, 16, 34, 100)):
            t.column_dimensions[col].width = w
        t.freeze_panes = "A2"

    out = a.out or (os.path.splitext(a.checklist)[0] + " - Status.xlsx")
    if os.path.abspath(out) == os.path.abspath(a.checklist):
        sys.exit("refusing to overwrite the source checklist")
    wb.save(out)
    print("saved %s  (%d rows written, %d left blank)"
          % (out, written, sum(1 for v in rows.values() if not v[0])))


if __name__ == "__main__":
    main()
